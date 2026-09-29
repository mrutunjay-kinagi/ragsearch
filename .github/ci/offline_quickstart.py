"""
Clean-install packaging check (#106), run by .github/workflows/ci.yml against the built wheel.

Runs outside the repository with only the installed wheel on the path, and:
1. checks that ragsearch comes from site-packages and that chromadb is NOT installed (#125);
2. runs the documented quickstart script (docs/quickstart.md) exactly as written, on its CSV;
3. runs setup() -> search() -> answer() on a generated multi-page PDF and a generated DOCX.

Providers are replaced by deterministic offline models and all network access is blocked, so
no real provider API is ever called.

Usage: python offline_quickstart.py <path to docs/quickstart.md>
"""

import contextlib
import importlib
import importlib.util
import io
import os
import re
import socket
import sys
import tempfile
from pathlib import Path


def _block_network():
    def refuse(*args, **kwargs):
        raise OSError("network access is blocked in the offline packaging check")

    socket.socket.connect = refuse
    socket.socket.connect_ex = refuse
    socket.create_connection = refuse


class OfflineEmbeddings:
    """Keyword-count embedding; one dimension per keyword plus a bias."""

    model_name = "offline-keywords"
    keywords = ["windshield", "glass", "denied", "harbour", "bollard", "mooring", "diagnosis", "radiculopathy"]

    class Response:
        def __init__(self, embeddings):
            self.embeddings = embeddings

    def embed(self, texts):
        return self.Response([[0.1] + [float(str(t).lower().count(k)) for k in self.keywords] for t in texts])


class OfflineLLM:
    def __init__(self):
        self.prompts = []

    def generate(self, prompt, **kwargs):
        self.prompts.append(prompt)
        return "Offline answer grounded in source [1]."


def write_text_pdf(path, pages):
    """Minimal multi-page PDF (Helvetica text), no extra dependencies."""
    def escape(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    for text in pages:
        lines = [text[i:i + 90] for i in range(0, len(text), 90)] or [""]
        ops = "BT /F1 10 Tf 12 TL 40 800 Td " + " ".join(f"({escape(line)}) Tj T*" for line in lines) + " ET"
        objects.append(f"<< /Length {len(ops.encode('latin-1'))} >>\nstream\n{ops}\nendstream")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objects)} 0 R >>"
        )
        page_ids.append(len(objects))
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{i} 0 R' for i in page_ids)}] /Count {len(page_ids)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1")
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode("latin-1")
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("latin-1")
    path.write_bytes(bytes(out))


def check(condition, message):
    if not condition:
        raise SystemExit(f"FAILED: {message}")
    print(f"ok: {message}")


def main(quickstart_md):
    import ragsearch

    check("site-packages" in ragsearch.__file__, f"ragsearch is the installed wheel ({ragsearch.__file__})")
    check(importlib.util.find_spec("chromadb") is None, "chromadb is not installed by default")

    _block_network()
    setup_module = importlib.import_module("ragsearch.setup")
    llm = OfflineLLM()
    setup_module.CohereClient = lambda *args, **kwargs: object()
    setup_module.create_embedding_model = lambda **kwargs: OfflineEmbeddings()
    setup_module.create_llm_client = lambda **kwargs: llm
    os.environ["COHERE_API_KEY"] = "offline-ci-key"

    with tempfile.TemporaryDirectory() as workdir:
        os.chdir(workdir)

        # 1. CSV: the documented quickstart, run as written.
        blocks = re.findall(r"^```python[^\n]*\n(.*?)^```", Path(quickstart_md).read_text(encoding="utf-8"), re.DOTALL | re.MULTILINE)
        check(len(blocks) == 1, "docs/quickstart.md has one python block")
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            exec(compile(blocks[0], "docs/quickstart.md", "exec"), {"__name__": "__main__"})
        printed = output.getvalue()
        check("Indexed 8 claims." in printed and "Answer: Offline answer" in printed, "quickstart runs on the CSV")

        from ragsearch import setup

        # 2. PDF: several pages, answer on the last page.
        pdf = Path(workdir) / "report.pdf"
        filler = "The committee reviewed access control policy and authentication records for the quarter. " * 25
        write_text_pdf(pdf, [f"Page {n}. {filler}" for n in range(1, 11)] + ["A cracked mooring bollard was found at the harbour."])
        engine = setup(pdf, os.environ["COHERE_API_KEY"], embeddings_dir="pdf-cache")
        top = engine.search("mooring bollard harbour", top_k=1)[0]
        answer = engine.answer("Where was the cracked mooring bollard found?", top_k=3)
        check(len(engine.index_data) > 1, f"PDF parsed and chunked ({len(engine.index_data)} chunks)")
        check("mooring bollard" in engine._result_text(top["metadata"]), "PDF late-page content retrieved")
        check(answer["citations"] and answer["answer"].startswith("Offline answer"), "PDF answer with citations")

        # 3. DOCX.
        import docx

        document = docx.Document()
        document.add_paragraph("Patient intake summary.")
        document.add_paragraph("The recorded diagnosis is lumbar radiculopathy after a lifting injury.")
        docx_path = Path(workdir) / "intake.docx"
        document.save(str(docx_path))
        engine = setup(docx_path, os.environ["COHERE_API_KEY"], embeddings_dir="docx-cache")
        answer = engine.answer("What is the diagnosis?", top_k=1)
        check("radiculopathy" in answer["context"], "DOCX parsed and its text sent as context")
        check(answer["citations"], "DOCX answer with citations")

    print("offline packaging check passed")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
